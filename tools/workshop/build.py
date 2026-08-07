#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Workshop-Generator – erzeugt aus einer JSON-Content-Spec die drei gebrandeten
Bausteine eines WMDG-Workshops:

  * WMDG-Workshop-<Stem>.pptx          – editierbare Präsentation (mit Notizen)
  * WMDG-Workbook-<Stem>.pdf           – Teilnehmer-Workbook (A4)
  * WMDG-Moderationsplan-<Stem>.pdf    – Moderations-/Ablaufplan (A4)

Design (Navy/Leaf/Teal, nummeriertes Motiv, Gehirn-Logo) ist identisch zu den
vier bestehenden Workshops. Aufruf:

    python3 tools/workshop/build.py <spec.json> [<spec2.json> ...]

Die Dateien landen in docs/workshop/<slug>/ und werden zusätzlich nach
content/vorlagen/workshop/ gespiegelt (Admin-Galerie). Danach:

    npm run vorlagen:galerie

Abhängigkeiten: python-pptx, Pillow, playwright (nutzt vorinstalliertes Chromium).
"""

import base64
import html as _html
import json
import os
import sys

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.environ.get("REPO_ROOT", os.path.abspath(os.path.join(HERE, "..", "..")))
ASSETS = os.path.join(HERE, "assets")
PDF_ASSETS = os.path.join(ROOT, "tools", "pdf", "assets")

DOCS_DIR = os.path.join(ROOT, "docs", "workshop")
MIRROR_DIR = os.path.join(ROOT, "content", "vorlagen", "workshop")

# --------------------------------------------------------------------------- #
# Marke
# --------------------------------------------------------------------------- #
NAVY = RGBColor(0x08, 0x10, 0x2A)
PANEL = RGBColor(0x1A, 0x22, 0x33)
PANEL2 = RGBColor(0x12, 0x1B, 0x2E)
TEAL = RGBColor(0x21, 0xB2, 0xBD)
TEAL_BRIGHT = RGBColor(0x34, 0xC4, 0xC4)
LEAF = RGBColor(0x8C, 0xC6, 0x3F)
LEAF_BRIGHT = RGBColor(0xA3, 0xD6, 0x4F)
GREEN = RGBColor(0x4F, 0x9E, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xAE, 0xB6, 0xC6)
MUTED2 = RGBColor(0xD8, 0xDB, 0xE6)
MUTED3 = RGBColor(0xC3, 0xCC, 0xDD)
# Helle Inhaltsfolien (weißer Grund)
INK = RGBColor(0x1A, 0x22, 0x33)      # Überschriften auf Weiß
SOFT = RGBColor(0x48, 0x52, 0x4E)     # Fließtext auf Weiß
NUM_DARK = RGBColor(0x08, 0x10, 0x2A)  # Zahl auf dem Farbkreis
MINT = RGBColor(0xE7, 0xF4, 0xF5)     # Reflexions-Karte (hell-teal)
GREENTINT = RGBColor(0xF2, 0xF6, 0xEC)  # Zusammenfassungs-Karten (hell-grün)
DEEPNAVY = RGBColor(0x0F, 0x1E, 0x44)  # Preis-Karte
LIGHTONNAVY = RGBColor(0xE7, 0xEC, 0xF5)  # Fließtext auf Navy

HEAD = "Cambria"
BODY = "Calibri"

EMU_IN = 914400
PXW = 12192000  # 13.333in – 16:9
PXH = 6858000   # 7.5in


def IN(v):
    return Emu(int(v * EMU_IN))


# --------------------------------------------------------------------------- #
# PPTX-Hilfen
# --------------------------------------------------------------------------- #
def _set_tracking(run, pts):
    """Zeichenabstand (Letter-Spacing) auf einem Run setzen."""
    rPr = run.font._element  # bereits das rPr-Element (CT_TextCharacterProperties)
    rPr.set("spc", str(int(pts * 100)))


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def _no_autosize(tf):
    # Kein automatisches Anpassen der Form an den Text.
    from pptx.enum.text import MSO_AUTO_SIZE
    tf.auto_size = MSO_AUTO_SIZE.NONE


def textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(IN(l), IN(t), IN(w), IN(h))
    tf = tb.text_frame
    tf.word_wrap = True
    _no_autosize(tf)
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    return tf


def para(tf, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    return p


def run(p, text, size, color, bold=False, italic=False, font=BODY, tracking=None):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    if tracking is not None:
        _set_tracking(r, tracking)
    return r


def kicker(tf, text, color=LEAF, size=12.5, first=True):
    p = para(tf, first=first)
    p.space_after = Pt(6)
    run(p, text.upper(), size, color, bold=True, font=BODY, tracking=2.4)
    return p


def bg_solid(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def bg_image(slide, path):
    slide.shapes.add_picture(path, 0, 0, Emu(PXW), Emu(PXH))


def circle(slide, l, t, d, number, fill=TEAL, num_color=WHITE, size=16):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, IN(l), IN(t), IN(d), IN(d))
    _fill(sh, fill)
    tf = sh.text_frame
    tf.word_wrap = False
    _no_autosize(tf)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, str(number), size, num_color, bold=True, font=HEAD)
    return sh


def rrect(slide, l, t, w, h, color, radius=0.12):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(l), IN(t), IN(w), IN(h))
    _fill(sh, color)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    return sh


def hairline(slide, l, t, w, color=TEAL, weight=1.5):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, IN(l), IN(t), IN(w), IN(0.012))
    _fill(sh, color)
    return sh


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


# --------------------------------------------------------------------------- #
# Folien-Typen  –  Design 1:1 nach den bestehenden Decks („Deinen Kopf
# verstehen" als Vorbild): Hero-Folien auf Navy-Bild, Inhaltsfolien auf Weiß.
# --------------------------------------------------------------------------- #
ML = 0.8            # linker Rand
CW = 13.333 - 2 * ML  # Inhaltsbreite
CIRCLE = os.path.join(ASSETS, "circle.png")


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg_white(slide):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE


def disc(slide, l, t, d, number, size):
    """Nummern-Kreis: Farbverlauf-Scheibe (Bild) mit dunkler Zahl darüber."""
    slide.shapes.add_picture(CIRCLE, IN(l), IN(t), IN(d), IN(d))
    tf = textbox(slide, l, t, d, d, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, first=True)
    p.alignment = PP_ALIGN.CENTER
    run(p, str(number), size, NUM_DARK, bold=True, font=HEAD)


def head_light(slide, kick, title, title_size=30, title_w=11.7):
    """Grüner Kicker + dunkle Serifen-Überschrift auf Weiß."""
    tf = textbox(slide, ML, 0.7, 8.0, 0.35)
    kicker(tf, kick, GREEN, 12)
    tf = textbox(slide, ML, 1.05, title_w, 1.0)
    p = para(tf, first=True)
    p.line_spacing = 1.03
    run(p, title, title_size, INK, bold=True, font=HEAD)


# ------------------------------- Hero-Folien -------------------------------- #
def slide_title(prs, s):
    sl = new_slide(prs)
    bg_image(sl, os.path.join(ASSETS, "bg-title.png"))
    sl.shapes.add_picture(os.path.join(ASSETS, "brain.png"), IN(9.7), IN(1.85),
                          width=IN(3.0), height=IN(3.0))
    tf = textbox(sl, 0.85, 1.2, 8.0, 0.35)
    kicker(tf, s["eyebrow"], LEAF_BRIGHT, 12)
    tf = textbox(sl, 0.8, 1.8, 8.7, 2.5)
    p = para(tf, first=True)
    p.line_spacing = 1.0
    run(p, s["title1"], 48, WHITE, bold=True, font=HEAD)
    p2 = tf.add_paragraph()
    p2.line_spacing = 1.0
    run(p2, s["title2"], 44, TEAL_BRIGHT, bold=False, font=HEAD)
    tf = textbox(sl, 0.85, 4.55, 8.3, 1.0)
    p = para(tf, first=True)
    p.line_spacing = 1.2
    run(p, s["subtitle"], 16, MUTED2, font=BODY)
    tf = textbox(sl, 0.85, 6.5, 9.0, 0.4)
    p = para(tf, first=True)
    run(p, s["author"], 13, MUTED, bold=True, font=BODY)
    notes(sl, s.get("notesTitle") or "Begrüßung. Kurz vorstellen. Rahmen setzen: heute geht es ums bewusste Sehen, nicht um Ratschläge. Vertraulichkeit im Raum betonen.")


def slide_statement(prs, s):
    k = s["kernbotschaft"]
    sl = new_slide(prs)
    bg_image(sl, os.path.join(ASSETS, "bg-divider.png"))
    tf = textbox(sl, 1.1, 2.2, 11.1, 2.6, anchor=MSO_ANCHOR.TOP)
    p = para(tf, first=True)
    p.line_spacing = 1.12
    run(p, k["pre"], 40, WHITE, bold=True, font=HEAD)
    run(p, k["accent"], 40, LEAF_BRIGHT, bold=True, font=HEAD)
    run(p, k["post"], 40, WHITE, bold=True, font=HEAD)
    if k.get("tail"):
        tf = textbox(sl, 1.12, 5.0, 10.0, 0.9)
        p = para(tf, first=True)
        p.line_spacing = 1.2
        run(p, k["tail"], 17, MUTED3, font=BODY)
    notes(sl, k.get("notes") or "Wirken lassen. Kurze Stille nach dem Satz.")


def slide_divider(prs, s, d):
    sl = new_slide(prs)
    bg_image(sl, os.path.join(ASSETS, "bg-divider.png"))
    tf = textbox(sl, 1.1, 2.35, 8.0, 0.35)
    kicker(tf, d["kicker"], LEAF_BRIGHT, 12)
    tf = textbox(sl, 1.05, 2.75, 11.2, 1.2)
    p = para(tf, first=True)
    run(p, d["title"], 42, WHITE, bold=True, font=HEAD)
    tf = textbox(sl, 1.1, 4.05, 10.2, 0.9)
    p = para(tf, first=True)
    p.line_spacing = 1.2
    run(p, d["subtitle"], 17, MUTED3, font=BODY)
    notes(sl, d.get("notes") or "Übergang zum Kern. Betonen: eine Landkarte, kein starres Schema.")


def slide_offer(prs, s):
    a = s["angebot"]
    sl = new_slide(prs)
    bg_image(sl, os.path.join(ASSETS, "bg-title.png"))
    tf = textbox(sl, 0.85, 1.1, 8.0, 0.35)
    kicker(tf, a["kicker"], LEAF_BRIGHT, 12)
    tf = textbox(sl, 0.8, 1.5, 8.6, 1.0)
    p = para(tf, first=True)
    p.line_spacing = 1.03
    run(p, a["title"], 38, WHITE, bold=True, font=HEAD)
    tf = textbox(sl, 0.9, 2.85, 7.9, 2.6)
    for i, f in enumerate(a["features"]):
        p = para(tf, first=(i == 0))
        p.line_spacing = 1.15
        p.space_after = Pt(9)
        run(p, "›  ", 17, LEAF_BRIGHT, bold=True, font=BODY)
        run(p, f, 17, LIGHTONNAVY, font=BODY)
    # Preis-Karte
    rrect(sl, 9.3, 2.9, 3.3, 2.0, DEEPNAVY, radius=0.1)
    tf = textbox(sl, 9.3, 3.15, 3.3, 0.4)
    p = para(tf, first=True)
    p.alignment = PP_ALIGN.CENTER
    run(p, "Mitgliedschaft", 13, MUTED, font=BODY)
    tf = textbox(sl, 9.3, 3.55, 3.3, 0.8)
    p = para(tf, first=True)
    p.alignment = PP_ALIGN.CENTER
    run(p, a["preis"], 40, WHITE, bold=True, font=HEAD)
    run(p, "  " + a.get("preisSuffix", ""), 16, MUTED3, font=BODY)
    tf = textbox(sl, 9.1, 4.45, 3.7, 0.35)
    p = para(tf, first=True)
    p.alignment = PP_ALIGN.CENTER
    run(p, a["url"], 11.5, LEAF_BRIGHT, bold=True, font=BODY)
    tf = textbox(sl, 0.9, 5.7, 8.0, 0.4)
    p = para(tf, first=True)
    run(p, a["note"], 15, MUTED3, italic=True, font=BODY)
    notes(sl, a.get("notes") or "Angebot ruhig, ohne Druck. Einladung, nicht Verkauf. Auf den kostenlosen Bewusstseinstest hinweisen.")


def slide_closing(prs, s):
    c = s["abschluss"]
    sl = new_slide(prs)
    bg_image(sl, os.path.join(ASSETS, "bg-divider.png"))
    # Marken-Logo mittig über dem Abschlusstitel (Bild 640×588).
    logo_h = 1.7
    logo_w = logo_h * 640 / 588
    sl.shapes.add_picture(os.path.join(ASSETS, "brain.png"),
                          IN((13.333 - logo_w) / 2), IN(2.05),
                          width=IN(logo_w), height=IN(logo_h))
    tf = textbox(sl, 1.0, 4.1, 11.33, 1.0)
    p = para(tf, first=True)
    p.alignment = PP_ALIGN.CENTER
    run(p, c["pre"], 40, WHITE, bold=True, font=HEAD)
    run(p, c["accent"], 40, LEAF_BRIGHT, bold=True, font=HEAD)
    run(p, c["post"], 40, WHITE, bold=True, font=HEAD)
    tf = textbox(sl, 1.0, 5.2, 11.33, 0.5)
    p = para(tf, first=True)
    p.alignment = PP_ALIGN.CENTER
    run(p, c["thanks"], 18, MUTED3, font=BODY)
    tf = textbox(sl, 1.0, 6.4, 11.33, 0.4)
    p = para(tf, first=True)
    p.alignment = PP_ALIGN.CENTER
    run(p, c["url"], 14, TEAL_BRIGHT, bold=True, font=BODY)
    notes(sl, c.get("notes") or "Abschluss. Raum für Fragen. Danken. Zum Bewusstseinstest / zur Mitgliedschaft einladen.")


# ------------------------------ Inhaltsfolien ------------------------------- #
def slide_agenda(prs, s):
    a = s["agenda"]
    sl = new_slide(prs)
    bg_white(sl)
    head_light(sl, "ÜBERBLICK", "Was dich heute erwartet", 38, 11.0)
    top = 2.15
    step = 0.98
    for i, it in enumerate(a):
        y = top + i * step
        disc(sl, 0.85, y, 0.62, i + 1, 21)
        tf = textbox(sl, 1.75, y - 0.04, 10.6, 0.42)
        p = para(tf, first=True)
        run(p, it["title"], 18, INK, bold=True, font=BODY)
        tf = textbox(sl, 1.75, y + 0.36, 10.6, 0.42)
        p = para(tf, first=True)
        run(p, it["desc"], 14, SOFT, font=BODY)
    notes(sl, s.get("notesAgenda") or "Agenda ruhig durchgehen. Zeitrahmen nennen. Frage in die Runde: Was möchtet ihr heute mitnehmen?")


def slide_metric(prs, s):
    a = s["ausgangspunkt"]
    sl = new_slide(prs)
    bg_white(sl)
    tf = textbox(sl, ML, 0.7, 8.0, 0.35)
    kicker(tf, a["kicker"], GREEN, 12)
    tf = textbox(sl, ML, 1.05, 7.7, 1.3)
    p = para(tf, first=True)
    p.line_spacing = 1.05
    run(p, a["title"], 32, INK, bold=True, font=HEAD)
    tf = textbox(sl, ML, 2.5, 7.4, 2.6)
    p = para(tf, first=True)
    p.line_spacing = 1.3
    p.space_after = Pt(12)
    run(p, a["p1"], 16.5, SOFT, font=BODY)
    p2 = tf.add_paragraph()
    p2.line_spacing = 1.3
    run(p2, a["p2"], 16.5, SOFT, font=BODY)
    # Navy-Kennzahl-Karte rechts
    rrect(sl, 8.7, 1.7, 3.9, 4.2, NAVY, radius=0.08)
    y = 2.55
    if a.get("metricPre"):
        tf = textbox(sl, 8.9, 2.35, 3.5, 0.35)
        p = para(tf, first=True)
        p.alignment = PP_ALIGN.CENTER
        run(p, a["metricPre"].upper(), 12, LEAF_BRIGHT, bold=True, tracking=2)
        y = 2.8
    val = a["metric"]
    vsize = 40 if len(str(val)) <= 7 else 30
    tf = textbox(sl, 8.75, y, 3.8, 1.15)
    p = para(tf, first=True)
    p.alignment = PP_ALIGN.CENTER
    run(p, val, vsize, LEAF_BRIGHT, bold=True, font=HEAD)
    tf = textbox(sl, 8.9, 3.95, 3.5, 0.5)
    p = para(tf, first=True)
    p.alignment = PP_ALIGN.CENTER
    run(p, a["metricLabel"], 16, WHITE, bold=True, font=BODY)
    tf = textbox(sl, 9.0, 4.5, 3.3, 1.2)
    p = para(tf, first=True)
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.15
    run(p, a["metricNote"], 13, MUTED3, font=BODY)
    notes(sl, a.get("notes") or "Kennzahl verankern. Überleitung zur ersten Übung.")


def slide_exercise(prs, s, ex):
    sl = new_slide(prs)
    bg_white(sl)
    tf = textbox(sl, ML, 0.7, 11.0, 0.35)
    kicker(tf, ex["kicker"], GREEN, 12)
    tf = textbox(sl, ML, 1.05, 11.6, 0.9)
    p = para(tf, first=True)
    run(p, ex["title"], 36, INK, bold=True, font=HEAD)
    steps = ex["steps"]
    top = 2.35
    span = 3.55
    step = min(0.92, span / max(len(steps), 1)) if len(steps) > 4 else 0.92
    for i, st in enumerate(steps):
        y = top + i * step
        disc(sl, 0.85, y, 0.56, i + 1, 19)
        tf = textbox(sl, 1.7, y - 0.02, 7.0, step, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, first=True)
        p.line_spacing = 1.1
        run(p, st, 16.5, INK, font=BODY)
    # MITNEHMEN-Karte rechts (Navy)
    rrect(sl, 9.15, 2.35, 3.45, 3.5, NAVY, radius=0.09)
    tf = textbox(sl, 9.4, 2.65, 3.0, 0.3)
    kicker(tf, "MITNEHMEN", LEAF_BRIGHT, 11)
    tf = textbox(sl, 9.4, 3.05, 3.0, 2.6)
    p = para(tf, first=True)
    p.line_spacing = 1.2
    run(p, ex["mitnehmen"], 18, WHITE, font=BODY)
    notes(sl, ex.get("notes") or "Übung anleiten und Zeit lassen. Danach 2–3 Stimmen einsammeln. Nichts bewerten.")


def slide_overview(prs, s):
    o = s["landkarte"]
    items = o["items"]
    sl = new_slide(prs)
    bg_white(sl)
    head_light(sl, o["kicker"], o["title"], 28, 11.7)
    n = len(items)
    col = (n + 1) // 2
    top = 2.0
    span = 4.7
    rows_max = max(col, n - col)
    step = min(1.18, span / max(rows_max, 1))
    for i, it in enumerate(items):
        c = 0 if i < col else 1
        r = i if i < col else i - col
        xc = 0.8 if c == 0 else 7.05
        xt = 1.7 if c == 0 else 7.95
        y = top + r * step
        disc(sl, xc, y, 0.66, i + 1, 22)
        tf = textbox(sl, xt, y - 0.02, 5.2, 0.42)
        p = para(tf, first=True)
        run(p, it["title"], 15, INK, bold=True, font=BODY)
        if it.get("sub"):
            tf = textbox(sl, xt, y + 0.35, 5.25, 0.5)
            p = para(tf, first=True)
            run(p, it["sub"], 11.5, TEAL, font=BODY)
    notes(sl, o.get("notes") or "Überblick geben, noch nicht vertiefen. Ankündigen: Wir gehen jedes einzeln durch.")


def slide_module(prs, s, m):
    sl = new_slide(prs)
    bg_white(sl)
    num = m["badge"].split()[-1]
    disc(sl, 0.85, 0.75, 1.5, num, 51)
    tf = textbox(sl, 2.65, 0.7, 8.0, 0.4)
    kicker(tf, m["badge"], GREEN, 13)
    tf = textbox(sl, 2.6, 1.05, 9.9, 1.0)
    p = para(tf, first=True)
    p.line_spacing = 1.02
    run(p, m["title"], 31, INK, bold=True, font=HEAD)
    tf = textbox(sl, 2.62, 2.05, 9.9, 0.6)
    p = para(tf, first=True)
    run(p, m["sub"], 18, TEAL, bold=False, font=BODY)
    tf = textbox(sl, 0.85, 3.6, 8.0, 2.2)
    p = para(tf, first=True)
    p.line_spacing = 1.3
    run(p, m["body"], 17, SOFT, font=BODY)
    # REFLEXION-Karte (hell-teal)
    rrect(sl, 9.15, 3.5, 3.45, 2.8, MINT, radius=0.09)
    tf = textbox(sl, 9.4, 3.75, 3.0, 0.3)
    kicker(tf, "REFLEXION", TEAL, 11)
    tf = textbox(sl, 9.4, 4.15, 3.0, 2.0)
    p = para(tf, first=True)
    p.line_spacing = 1.18
    run(p, m["reflexion"], 16, INK, font=BODY)
    notes(sl, m.get("notes") or f"{m['badge']} – Kernaussage erklären, eigenes Beispiel erzählen. Reflexionsfrage in die Runde oder ins Workbook geben.")


def slide_summary(prs, s):
    z = s["zusammenfassung"]
    items = z["items"]
    sl = new_slide(prs)
    bg_white(sl)
    head_light(sl, z["kicker"], z["title"], 34, 11.6)
    n = len(items)
    gap = 0.35
    cw = (11.75 - (n - 1) * gap) / n
    top = 2.4
    for i, it in enumerate(items):
        x = 0.8 + i * (cw + gap)
        rrect(sl, x, top, cw, 3.4, GREENTINT, radius=0.07)
        disc(sl, x + 0.35, top + 0.35, 0.7, i + 1, 24)
        tf = textbox(sl, x + 0.35, top + 1.3, cw - 0.7, 0.5)
        p = para(tf, first=True)
        run(p, it["title"], 22, INK, bold=True, font=HEAD)
        tf = textbox(sl, x + 0.35, top + 1.85, cw - 0.7, 1.4)
        p = para(tf, first=True)
        p.line_spacing = 1.2
        run(p, it["text"], 14, SOFT, font=BODY)
    notes(sl, z.get("notes") or "Zusammenfassen. Der rote Faden. Überleitung zum Angebot.")


def build_pptx(s, out_path):
    prs = Presentation()
    prs.slide_width = Emu(PXW)
    prs.slide_height = Emu(PXH)

    slide_title(prs, s)
    slide_agenda(prs, s)
    slide_statement(prs, s)
    slide_metric(prs, s)
    slide_exercise(prs, s, s["uebung1"])
    slide_divider(prs, s, s["kapitelIntro"])
    slide_overview(prs, s)
    for m in s["module"]:
        slide_module(prs, s, m)
    slide_exercise(prs, s, s["uebung2"])
    slide_summary(prs, s)
    slide_offer(prs, s)
    slide_closing(prs, s)

    prs.save(out_path)
    return len(prs.slides._sldIdLst)


# --------------------------------------------------------------------------- #
# PDF-Hilfen (HTML -> Chromium)
# --------------------------------------------------------------------------- #
def _enc(path, mime):
    with open(path, "rb") as f:
        return "data:%s;base64,%s" % (mime, base64.b64encode(f.read()).decode())


def esc(x):
    return _html.escape(str(x), quote=False)


FONTS_CSS = open(os.path.join(PDF_ASSETS, "fonts.css")).read()
BRAIN_URI = _enc(os.path.join(ASSETS, "brain.png"), "image/png")

BASE_CSS = r"""
*{margin:0;padding:0;box-sizing:border-box;}
html{background:#fff;-webkit-print-color-adjust:exact;print-color-adjust:exact;}
body{font-family:'Inter',ui-sans-serif,system-ui,sans-serif;color:#16231f;}
.serif{font-family:'Fraunces',Georgia,serif;}
@page{size:A4;margin:16mm 17mm 16mm;}
:root{
  --navy:#08102a; --ink:#16231f; --soft:#48524e;
  --teal:#21b2bd; --teal-b:#34c4c4; --leaf:#8cc63f; --green:#4f9e1c;
  --surface:#f6f4ee; --hair:#e4ded0; --muted:#aeb6c6;
}
.kicker{font-size:10.5px;font-weight:700;letter-spacing:.18em;text-transform:uppercase;color:var(--green);}
.kicker.teal{color:var(--teal);}
h1.doc{font-family:'Fraunces',serif;font-weight:600;font-size:30px;color:var(--navy);margin-top:4px;line-height:1.12;}
h2.doc{font-family:'Fraunces',serif;font-weight:600;font-size:21px;color:var(--navy);}
p.body{font-size:13.5px;line-height:1.62;color:#2c3a35;margin-top:10px;}
"""


def _pdf_shell(title, css, body):
    return f"""<!doctype html><html lang="de"><head><meta charset="utf-8">
<title>{esc(title)}</title><style>{FONTS_CSS}\n{BASE_CSS}\n{css}</style></head>
<body>{body}</body></html>"""


# ------------------------------ Workbook ----------------------------------- #
WB_CSS = r"""
@page cover{margin:0;}
.cover{page:cover;position:relative;width:210mm;height:297mm;
  padding:0 24mm;display:flex;flex-direction:column;align-items:center;justify-content:center;
  text-align:center;color:#eaf0ff;overflow:hidden;
  background:
    radial-gradient(120% 80% at 82% 6%, rgba(52,196,196,.20), transparent 55%),
    radial-gradient(90% 60% at 12% 98%, rgba(140,198,63,.14), transparent 55%),
    linear-gradient(158deg,#0a1330,#0a1024);}
.cover .eye{font-size:13px;font-weight:700;letter-spacing:.34em;color:var(--leaf);text-transform:uppercase;margin-bottom:26px;}
.cover img{width:150px;height:auto;margin-bottom:30px;filter:drop-shadow(0 8px 30px rgba(52,196,196,.35));}
.cover h1{font-family:'Fraunces',serif;font-weight:600;font-size:38px;line-height:1.18;color:#fff;}
.cover h1 .ac{color:var(--teal-b);}
.cover .sub{margin-top:16px;font-size:15px;line-height:1.55;color:#c6cfe6;max-width:78%;}
.cover .name{position:absolute;bottom:20mm;left:0;right:0;font-size:12.5px;color:#aeb6c6;font-weight:600;}
.cover .name b{color:#eaf0ff;font-weight:600;}
.page{page-break-before:always;}
.intro{margin-top:2px;}
.callout{margin-top:18px;border-radius:16px;padding:18px 22px;color:#eaf0ff;
  background:radial-gradient(70% 130% at 88% 0%, rgba(52,196,196,.26), transparent 60%),
             linear-gradient(140deg,#08102a,#12244d);}
.callout .k{font-size:10px;font-weight:700;letter-spacing:.18em;text-transform:uppercase;color:var(--leaf);}
.callout .q{font-family:'Fraunces',serif;font-style:italic;font-size:14.5px;line-height:1.5;margin-top:8px;color:#eef2ff;}
.steps{margin-top:12px;list-style:none;counter-reset:st;}
.steps li{font-size:13.5px;line-height:1.5;color:#26332e;margin:7px 0;padding-left:4px;}
.steps li b{color:var(--navy);}
.prompt{margin-top:16px;}
.lines{margin-top:8px;}
.lines .ln{border-bottom:1.4px solid #d9d3c4;height:30px;}
.sec{page-break-inside:avoid;margin-top:26px;}
.sec:first-of-type{margin-top:6px;}
"""


def workbook_html(s):
    wb = s["workbook"]
    cover = f"""<section class="cover">
      <div class="eye">Workshop-Begleitheft</div>
      <img src="{BRAIN_URI}" alt="">
      <h1>{esc(s['title1'])} <span class="ac">{esc(s['title2'])}</span></h1>
      <div class="sub">{esc(wb.get('coverSubtitle',''))}</div>
      <div class="name"><b>Name:</b> ______________________ · werdemeisterdeinergedanken.de</div>
    </section>"""

    intro = wb["intro"]
    parts = [f"""<div class="page"><div class="kicker">So arbeitest du mit diesem Heft</div>
      <h1 class="doc">Willkommen</h1>
      <p class="body">{esc(intro['body'])}</p>
      <div class="callout"><div class="k">{esc(intro.get('grundKicker','Die Grundübung'))}</div>
        <div class="q">{esc(intro['grundText'])}</div></div>"""]

    # Sektionen (Module) – erste direkt auf der Intro-Seite, Rest je nach Umbruch
    for i, sec in enumerate(wb["sections"]):
        steps = "".join(f"<li><b>{j+1}.</b> {esc(st)}</li>" for j, st in enumerate(sec.get("steps", [])))
        lines = "".join('<div class="ln"></div>' for _ in range(int(sec.get("lines", 3))))
        parts.append(f"""<div class="sec">
          <div class="kicker teal">{esc(sec['kicker'])}</div>
          <h2 class="doc">{esc(sec['title'])}</h2>
          <p class="body">{esc(sec['body'])}</p>
          {f'<div class="kicker teal" style="margin-top:14px">{esc(sec.get("uebungKicker",""))}</div>' if sec.get('uebungKicker') else ''}
          {f'<ul class="steps">{steps}</ul>' if steps else ''}
          <div class="prompt kicker teal">{esc(sec['prompt'])}</div>
          <div class="lines">{lines}</div>
        </div>""")
    parts.append("</div>")
    body = cover + "".join(parts)
    return _pdf_shell(f"WMDG Workbook – {s['title1']} {s['title2']}", WB_CSS, body)


# --------------------------- Moderationsplan -------------------------------- #
MOD_CSS = r"""
.head{display:flex;gap:14px;align-items:flex-start;border-bottom:2px solid var(--navy);padding-bottom:14px;}
.head img{width:56px;height:auto;margin-top:2px;}
.head .k{font-size:10.5px;font-weight:700;letter-spacing:.2em;text-transform:uppercase;color:var(--green);}
.head h1{font-family:'Fraunces',serif;font-weight:600;font-size:25px;color:var(--navy);line-height:1.14;margin-top:3px;}
.head h1 .ac{color:var(--teal);}
.pills{display:flex;flex-wrap:wrap;gap:12px;margin:16px 0 4px;}
.pill{background:var(--navy);color:#eaf0ff;border-radius:999px;padding:8px 16px;font-size:12px;}
.pill b{color:var(--leaf);font-weight:700;}
.cards{display:flex;gap:16px;margin-top:16px;}
.card{flex:1;background:var(--surface);border:1px solid var(--hair);border-radius:14px;padding:16px 18px;}
.card h3{font-family:'Fraunces',serif;font-weight:600;font-size:15px;color:var(--navy);margin-bottom:8px;}
.card ul{list-style:none;}
.card li{font-size:12px;line-height:1.5;color:#3a453f;padding-left:14px;position:relative;margin:5px 0;}
.card li:before{content:"•";position:absolute;left:0;color:var(--green);font-weight:700;}
h2.ablauf{font-family:'Fraunces',serif;font-weight:600;font-size:19px;color:var(--navy);margin:22px 0 10px;}
table{width:100%;border-collapse:collapse;}
thead th{background:var(--navy);color:#eaf0ff;font-size:9.5px;letter-spacing:.12em;text-transform:uppercase;
  text-align:left;padding:9px 10px;font-weight:700;}
thead th:first-child{border-radius:8px 0 0 8px;}
thead th:last-child{border-radius:0 8px 8px 0;}
tbody td{font-size:11.5px;line-height:1.4;padding:11px 10px;border-bottom:1px solid #eceadf;vertical-align:top;color:#2f3a35;}
tbody tr:nth-child(even){background:#faf9f4;}
.t-zeit{color:var(--teal);font-weight:700;white-space:nowrap;}
.t-dur{color:#8a938d;white-space:nowrap;}
.t-block{font-family:'Fraunces',serif;font-weight:600;color:var(--navy);}
.pause td{color:var(--green);font-style:italic;}
.pause .t-block{color:var(--green);}
.foot{margin-top:18px;font-size:10.5px;color:#8a938d;}
"""


def moderation_html(s):
    m = s["moderation"]
    ziel = "".join(f"<li>{esc(x)}</li>" for x in m["ziel"])
    mat = "".join(f"<li>{esc(x)}</li>" for x in m["material"])
    rows = ""
    for r in m["ablauf"]:
        cls = ' class="pause"' if r.get("pause") else ""
        rows += f"""<tr{cls}><td class="t-zeit">{esc(r['time'])}</td>
          <td class="t-dur">{esc(r['dur'])}</td>
          <td class="t-block">{esc(r['block'])}</td>
          <td>{esc(r['what'])}</td>
          <td>{esc(r.get('material','—'))}</td></tr>"""
    body = f"""
    <div class="head"><img src="{BRAIN_URI}" alt="">
      <div><div class="k">Moderations- &amp; Ablaufplan</div>
        <h1>Workshop: {esc(s['title1'])} <span class="ac">{esc(s['title2'])}</span></h1></div>
    </div>
    <div class="pills">
      <div class="pill"><b>Dauer:</b> {esc(s['dauer'])}</div>
      <div class="pill"><b>Gruppe:</b> {esc(s['gruppe'])}</div>
      <div class="pill"><b>Format:</b> {esc('Präsenz oder online')}</div>
      <div class="pill"><b>Leitung:</b> {esc(s['leitung'])}</div>
    </div>
    <div class="cards">
      <div class="card"><h3>Ziel des Workshops</h3><ul>{ziel}</ul></div>
      <div class="card"><h3>Material-Checkliste</h3><ul>{mat}</ul></div>
    </div>
    <h2 class="ablauf">Ablauf</h2>
    <table><thead><tr><th>Zeit</th><th>Dauer</th><th>Block</th><th>Was passiert</th><th>Folien / Material</th></tr></thead>
    <tbody>{rows}</tbody></table>
    <div class="foot">Zeiten sind Richtwerte – passe sie an Gruppengröße und Tempo an. Preis auf der Angebots-Folie (49 €/Monat) vor dem Einsatz prüfen.</div>
    """
    return _pdf_shell(f"WMDG Moderationsplan – {s['title1']} {s['title2']}", MOD_CSS, body)


# --------------------------------------------------------------------------- #
def render_pdfs(html_jobs):
    """html_jobs: list of (html_string, out_path)."""
    from playwright.sync_api import sync_playwright
    exe = None
    for cand in (
        "/opt/pw-browsers/chromium-1194/chrome-linux/chrome",
    ):
        if os.path.exists(cand):
            exe = cand
            break
    with sync_playwright() as p:
        kw = {"executable_path": exe} if exe else {}
        b = p.chromium.launch(**kw)
        pg = b.new_page()
        for html, out in html_jobs:
            pg.set_content(html, wait_until="networkidle")
            pg.pdf(path=out, format="A4", print_background=True,
                   prefer_css_page_size=True)
        b.close()


def build_one(spec_path):
    s = json.load(open(spec_path, encoding="utf-8"))
    stem = s["fileStem"]
    slug = s["slug"]
    out_dir = os.path.join(DOCS_DIR, slug)
    os.makedirs(out_dir, exist_ok=True)

    pptx_path = os.path.join(out_dir, f"WMDG-Workshop-{stem}.pptx")
    wb_path = os.path.join(out_dir, f"WMDG-Workbook-{stem}.pdf")
    mod_path = os.path.join(out_dir, f"WMDG-Moderationsplan-{stem}.pdf")

    n = build_pptx(s, pptx_path)
    render_pdfs([
        (workbook_html(s), wb_path),
        (moderation_html(s), mod_path),
    ])

    # nach content/vorlagen/workshop/ spiegeln
    os.makedirs(MIRROR_DIR, exist_ok=True)
    import shutil
    for pth in (pptx_path, wb_path, mod_path):
        shutil.copy2(pth, os.path.join(MIRROR_DIR, os.path.basename(pth)))

    print(f"✓ {slug}: {n} Folien · {len(s['module'])} Module")
    print(f"    {os.path.relpath(pptx_path, ROOT)}")
    print(f"    {os.path.relpath(wb_path, ROOT)}")
    print(f"    {os.path.relpath(mod_path, ROOT)}")
    return {"slug": slug, "stem": stem, "slides": n}


def main(argv):
    if not argv:
        print("Aufruf: python3 tools/workshop/build.py <spec.json> [...]")
        return 1
    for spec in argv:
        build_one(spec)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
